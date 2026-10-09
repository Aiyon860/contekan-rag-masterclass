"""Ekstrak teks dari berbagai format. Hasil: list segmen (teks, nomor halaman/None)."""
import csv
from pathlib import Path

ALLOWED_EXT = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv"}
ROWS_PER_GROUP = 15


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _pdf(path: Path) -> list[tuple[str, int | None]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [((p.extract_text() or "").strip(), i + 1) for i, p in enumerate(reader.pages)]


def _cell_text(cell) -> str:
    return " ".join(cell.text.split())


def _unique_cells(row) -> list[str]:
    """Teks sel per baris; sel yang digabung (merged) hanya dihitung sekali."""
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(_cell_text(c))
    return out


def _table_block(table, heading: str | None) -> str:
    rows = [r for r in (_unique_cells(r) for r in table.rows) if any(r)]
    if not rows:
        return ""
    title = f"## {heading}\n" if heading else ""
    width = max(len(r) for r in rows)
    if width == 2:
        # tabel kunci-nilai: tiap baris jadi "Kunci: Nilai" supaya nilai selalu menempel pada kuncinya
        return title + "\n".join(f"{k}: {v}" for k, v in rows)
    # tabel biasa: baris pertama dianggap header
    header, body = rows[0], rows[1:]
    lines = ["Kolom: " + ", ".join(header)]
    for r in body:
        lines.append("; ".join(f"{h}: {v}" for h, v in zip(header, r) if v))
    return title + "\n".join(lines)


def _docx(path: Path) -> list[tuple[str, int | None]]:
    """Baca .docx sesuai urutan asli (paragraf dan tabel berselang-seling).

    Judul/paragraf sebelum tabel dipakai sebagai label tabel itu, jadi "AHMAD AIDIN" tidak terpisah
    dari tabel "Total Upah" miliknya.
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[str] = []
    heading: str | None = None  # judul/paragraf terakhir yang belum "dipakai" tabel
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = " ".join(para.text.split())
            if not text:
                continue
            if heading:  # paragraf sebelumnya bukan label tabel, simpan sebagai teks biasa
                blocks.append(heading)
            heading = text
        elif tag == "tbl":
            block = _table_block(Table(child, doc), heading)
            heading = None
            if block:
                blocks.append(block)
    if heading:
        blocks.append(heading)
    return [("\n\n".join(blocks), None)]


def _rows_to_groups(rows: list[list[str]], label: str | None = None) -> list[tuple[str, int | None]]:
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    prefix = f"[{label}] " if label else ""
    head_line = f"{prefix}Kolom: " + ", ".join(header)
    groups = []
    for i in range(0, len(body), ROWS_PER_GROUP):
        lines = []
        for r in body[i : i + ROWS_PER_GROUP]:
            lines.append("; ".join(f"{h}: {v}" for h, v in zip(header, r) if str(v).strip()))
        groups.append((head_line + "\n" + "\n".join(lines), None))
    return groups or [(head_line, None)]


def _csv(path: Path) -> list[tuple[str, int | None]]:
    with path.open(newline="", encoding="utf-8", errors="replace") as f:
        rows = [r for r in csv.reader(f) if any(c.strip() for c in r)]
    return _rows_to_groups(rows)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_sheets(path: Path, max_rows: int | None = None) -> list[tuple[str, list[list[str]]]]:
    """Baca Excel berdasarkan isi file (magic bytes), bukan ekstensi.

    Banyak aplikasi (CRM, ERP) menamai file Excel 97-2003 biner dengan akhiran .xlsx.
    """
    with path.open("rb") as f:
        magic = f.read(4)
    sheets: list[tuple[str, list[list[str]]]] = []
    if magic == b"PK\x03\x04":  # .xlsx (zip)
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in r):
                    rows.append([_cell(c) for c in r])
                if max_rows and len(rows) >= max_rows:
                    break
            sheets.append((ws.title, rows))
    elif magic == b"\xd0\xcf\x11\xe0":  # .xls biner (OLE2)
        import xlrd

        wb = xlrd.open_workbook(str(path))
        for ws in wb.sheets():
            rows = []
            for ri in range(ws.nrows):
                vals = []
                for ci, cell in enumerate(ws.row(ri)):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            vals.append(xlrd.xldate_as_datetime(cell.value, wb.datemode).isoformat(sep=" "))
                            continue
                        except Exception:
                            pass
                    vals.append("" if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else _cell(cell.value))
                if any(v.strip() for v in vals):
                    rows.append(vals)
                if max_rows and len(rows) >= max_rows:
                    break
            sheets.append((ws.name, rows))
    else:
        raise ValueError("Isi file tidak cocok dengan format Excel. Simpan ulang sebagai .xlsx atau .csv, lalu upload lagi.")
    return sheets


def _xlsx(path: Path) -> list[tuple[str, int | None]]:
    out = []
    for title, rows in read_sheets(path):
        out += _rows_to_groups(rows, label=title)
    return out


def parse_file(path: Path) -> list[tuple[str, int | None]]:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        segments = [(_read_text(path), None)]
    elif ext == ".pdf":
        segments = _pdf(path)
    elif ext == ".docx":
        segments = _docx(path)
    elif ext == ".csv":
        segments = _csv(path)
    elif ext == ".xlsx":
        segments = _xlsx(path)
    else:
        raise ValueError(f"Format {ext} tidak didukung")
    segments = [(t, p) for t, p in segments if t.strip()]
    if not segments:
        raise ValueError("Tidak ada teks yang bisa dibaca. Kalau ini PDF hasil scan, teksnya belum bisa diekstrak.")
    return segments
